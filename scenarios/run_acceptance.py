"""A synthetic multi-software research project, with independent numerical/content oracles.

This suite may create a hidden Office automation instance and execute code in installed
research engines. It never logs in, changes licenses, sends messages or modifies source data.
"""
import argparse
import csv
import importlib.util
import json
import math
import os
import subprocess
import sys
import time
import uuid
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from research_bridge.catalog import CATALOG, discover, load_config
from research_bridge.core import compile_job, run_compiled, save_json
from research_bridge.verify import verify_directory

def load_module(name,path):
    spec=importlib.util.spec_from_file_location(name,path)
    module=importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    return module

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("workspace",type=Path)
    parser.add_argument("--config")
    parser.add_argument("--only",help="Comma-separated engines, or all")
    parser.add_argument("--spss-timeout",type=int,default=35)
    parser.add_argument("--zotero-fixture",action="store_true",help="Launch an isolated synthetic Zotero library and validate real API reads")
    args=parser.parse_args()
    root=args.workspace.resolve(); root.mkdir(parents=True,exist_ok=True)
    selected=set(args.only.split(',')) if args.only else None
    if selected and not selected.issubset(CATALOG):
        parser.error("unknown --only engines: "+", ".join(sorted(selected-set(CATALOG))))
    cfg=load_config(args.config)
    prepared=load_module("statistics_fixture",Path(__file__).with_name("statistics")/"prepare.py").prepare_scenario(root/"fixtures/statistics")
    fixture=root/"fixtures"; fixture.mkdir(exist_ok=True)
    oracle=json.loads(Path(prepared["oracle"]).read_text())
    validators=load_module("statistics_oracles",Path(__file__).with_name("statistics")/"validate.py")
    checks=[]
    outputs={}

    def run(job, validator=None):
        engine=job["engine"]
        if selected is not None and engine not in selected: return None
        directory=root/"runs"/(job["id"]+"-"+uuid.uuid4().hex[:8])
        started=time.monotonic()
        try:
            plan=compile_job(job,directory,cfg)
            result=run_compiled(plan,max_timeout_s=max(180,args.spss_timeout))
            check=dict(id=job["id"],engine=engine,status=result["status"],scope=result.get("scope",plan["mode"]),application_verified=result.get("application_verified",False),workspace=str(directory))
            if result["status"]=="passed" and validator:
                check["oracle"]=validator(directory)
            elif result["status"]!="passed":
                check["reason"]=result.get("reason") or result.get("errors")
                check["process_diagnostics"]=[{k:p.get(k) for k in ("exit_code","timed_out","stderr")} for p in result.get("processes",[])]
            outputs[job["id"]]=directory
        except Exception as e:
            check=dict(id=job["id"],engine=engine,status="failed",reason=type(e).__name__+": "+str(e),workspace=str(directory),application_verified=False)
        check["duration_s"]=round(time.monotonic()-started,3)
        checks.append(check)
        save_json(root/"acceptance.json",dict(status="running",checks=checks))
        print(json.dumps({k:check[k] for k in ("id","status","duration_s")}),flush=True)
        return check

    def job(engine,op,id,source=None,params=None,**inputs):
        if source: inputs["source"]=str(source)
        return dict(schema_version="research.job.v1",id=id,engine=engine,operation=op,inputs=inputs,parameters=params or {},timeout_s=120)

    def numerical(directory):
        with (directory/"coefficients.csv").open(encoding="utf-8-sig",newline="") as f:
            rows=list(csv.DictReader(f))
        estimates={r["term"]:float(r["estimate"]) for r in rows}
        errors=[abs(estimates[k]-v) for k,v in oracle["coefficients"].items()]
        if max(errors)>1e-7: raise AssertionError("OLS coefficient oracle mismatch")
        metrics=json.loads((directory/"metrics.json").read_text())
        if metrics["n_complete"]!=535 or metrics["n_excluded"]!=41: raise AssertionError("missing-data audit mismatch")
        return dict(max_coefficient_abs_error=max(errors),n_complete=535,n_excluded=41)

    jobs=json.loads(Path(prepared["jobs"]).read_text())
    probe=job("spss","probe","spss-real-engine-probe")
    probe["timeout_s"]=args.spss_timeout
    probe_result=run(probe)
    for item in jobs:
        if item["engine"]=="spss" and (not probe_result or probe_result["status"]!="passed"):
            if selected is None or "spss" in selected:
                checks.append(dict(id=item["id"],engine="spss",status="blocked",scope="native",application_verified=False,reason="SPSS sentinel probe did not pass; analysis not run"))
            continue
        validator=(lambda directory:validators.check_ols(directory,oracle)) if item["operation"]=="analysis" and item["engine"]!="sas" else None
        if item["id"]=="longitudinal_mixed_r":
            validator=lambda directory:validators.check_mixed(directory,Path(prepared["data"]))
        run(item,validator)

    refs=[{"id":"r2026","title":"R: A Language and Environment for Statistical Computing","author":[{"literal":"R Core Team"}],"year":"2026","container-title":"R Foundation"},
          {"id":"study2026","title":"Synthetic intervention study: reproducibility fixture","author":[{"family":"Bridge","given":"Research"}],"year":"2026","DOI":"10.0000/synthetic-fixture"}]
    save_json(fixture/"references.json",refs)
    run(job("jabref","bibliography","bibtex-interchange",fixture/"references.json"))
    run(job("endnote","bibliography","ris-interchange",fixture/"references.json"))
    codes={"codes":[{"name":"Participation / 参与","description":"Participation mechanisms, with positive and negative cases.","color":"#245080","children":[{"name":"Institutional support / 制度支持","description":"Explicit mention of institutional help."},{"name":"Peer support / 同伴支持","description":"Explicit mention of peer help."}]},{"name":"Barriers / 阻碍","description":"Barriers to sustained participation","color":"#AA4020","children":[{"name":"Time constraints / 时间限制"},{"name":"Resource constraints / 资源限制"}]}]}
    save_json(fixture/"codebook.json",codes)
    def qdc(directory):
        tree=ET.parse(directory/"codebook.qdc")
        elements=tree.findall(".//{urn:QDA-XML:codebook:1.0}Code")
        if len(elements)!=6 or len({x.get("guid") for x in elements})!=6: raise AssertionError("hierarchical codebook structure mismatch")
        return dict(codes=6,consumer_application_verified=False)
    for engine in ("nvivo","maxqda","atlas-ti"):
        run(job(engine,"codebook",engine+"-codebook",fixture/"codebook.json"),qdc)
    for engine in ("prism","jasp"):
        run(job(engine,"table",engine+"-csv-interchange",Path(prepared["data"])))
    run(job("rstudio","project","rstudio-project-interchange"))

    # A real research chart source is prepared from the synthetic study, then exported by Inkscape.
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        rows=list(csv.DictReader(Path(prepared["data"]).open()))
        fig,ax=plt.subplots(figsize=(7,4))
        for group in ("control","active_a","active_b"):
            values=[]
            for visit in range(4):
                series=[float(r["score"]) for r in rows if r["group"]==group and int(r["time"])==visit and r["score"]]
                values.append(sum(series)/len(series))
            ax.plot(range(4),values,marker="o",label=group)
        ax.set(xlabel="Follow-up visit",ylabel="Mean synthetic score",title="Three-arm longitudinal intervention"); ax.legend(); fig.tight_layout(); fig.savefig(fixture/"study.svg"); plt.close(fig)
    except ImportError:
        # A missing scientific Python dependency is explicit, never silently replaced by a fake chart.
        checks.append(dict(id="matplotlib-source",engine="python",status="blocked",reason="matplotlib not installed",application_verified=False))
    if (fixture/"study.svg").exists():
        def png(directory):
            from PIL import Image
            im=Image.open(directory/"figure.png")
            if min(im.size)<1000 or len(im.convert("RGB").getcolors(im.width*im.height) or [])<50: raise AssertionError("raster export is empty or below publication resolution")
            return dict(width=im.width,height=im.height,dpi=300)
        run(job("inkscape","export","inkscape-publication-png",fixture/"study.svg",{"format":"png","dpi":300}),png)
        run(job("inkscape","export","inkscape-publication-pdf",fixture/"study.svg",{"format":"pdf"}))
    dot='digraph Study { rankdir=LR; node [shape=box]; Recruitment -> Randomization; Randomization -> Control; Randomization -> ActiveA; Randomization -> ActiveB; Control -> Followup; ActiveA -> Followup; ActiveB -> Followup; Followup -> Model; }'
    (fixture/"study.dot").write_text(dot)
    def dag(directory):
        svg=ET.parse(directory/"diagram.svg")
        groups=[e for e in svg.iter() if e.tag.endswith('g')]
        nodes=sum(e.get('class')=='node' for e in groups); edges=sum(e.get('class')=='edge' for e in groups)
        if nodes!=7 or edges!=8: raise AssertionError("Graphviz DAG node/edge oracle mismatch")
        return dict(nodes=nodes,edges=edges)
    run(job("graphviz","render","graphviz-study-dag",fixture/"study.dot",{"format":"svg"}),dag)
    (fixture/"origin-analysis.ogs").write_text("newbook; col(A)={1,2,3,4}; col(B)={2,4,7,9}; plotxy (1,2) plot:=200;\n",encoding="utf-8")
    run(job("origin","script","origin-native-project",fixture/"origin-analysis.ogs"))
    macro='''setBatchMode(true);
newImage("Synthetic calibration", "8-bit black", 128, 128, 1);
setColor(255);
for (i=0; i<6; i++) { makeRectangle(8+38*(i%3), 8+54*floor(i/3), 10+2*(i%3), 12+2*(i%3)); run("Fill"); }
makeRectangle(1,1,2,2); run("Fill");
run("Set Scale...", "distance=2 known=1 unit=um");
run("Set Measurements...", "area mean centroid perimeter redirect=None decimal=6");
run("Select None");
setThreshold(128,255);
run("Analyze Particles...", "size=4-Infinity display clear");
if (nResults!=6) exit("Synthetic particle count oracle failed");
saveAs("Results", "{{OUTPUT_DIR}}/measurements.csv");
'''
    (fixture/"particles.ijm").write_text(macro)
    def particles(directory):
        rows=list(csv.DictReader((directory/"measurements.csv").open()))
        area=sum(float(r["Area"]) for r in rows)
        if len(rows)!=6 or abs(area-256)>1e-6: raise AssertionError("ImageJ calibration/count/area oracle mismatch: "+str((len(rows),area)))
        return dict(particles=len(rows),total_area_um2=area,excluded_small_noise=True)
    run(job("imagej","macro","imagej-particle-calibration",fixture/"particles.ijm"),particles)

    md='''---
title: "科研桥接验收：模拟纵向干预研究"
author: "Synthetic Research Fixture"
---

# Study and methods

This synthetic study contains **144 subjects**, **576 observations**, and **41 missing outcomes**. No participant records are used.

| Group | Subjects | Visits |
|:--|--:|--:|
| Control | 48 | 4 |
| Active A | 48 | 4 |
| Active B | 48 | 4 |

The model is $y = X\\beta + \\epsilon$. Reproducible computation uses R [@r2026].

# 中文结论与限制

完整病例为 535 行。缺失值审计、多重比较和受试者层级重复抽样均保留原始证据。

## References
'''
    (fixture/"report.md").write_text(md,encoding="utf-8")
    bib=outputs.get("bibtex-interchange")
    extras={"bibliography":str(bib/"references.bib")} if bib else {}
    def docx(directory):
        with zipfile.ZipFile(directory/"report.docx") as z:
            xml=z.read("word/document.xml").decode()
        tree=ET.fromstring(xml)
        text="".join(tree.itertext())
        if "535" not in text or "中文结论" not in text or not tree.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl") or "R Core Team" not in text: raise AssertionError("DOCX text/table/citation oracle mismatch")
        return dict(unicode=True,table=True,citation=True,complete_cases=535)
    run(job("pandoc","convert","pandoc-research-docx",fixture/"report.md",{"format":"docx"},**extras),docx)
    qmd='''---
title: "Executable research report"
format: html
execute:
  echo: true
---

## Auditable results

```{r}
data <- read.csv("statistics/longitudinal.csv")
cat("OBSERVATIONS=", nrow(data), " COMPLETE=", sum(!is.na(data$score)), sep="")
stopifnot(nrow(data)==576, sum(!is.na(data$score))==535)
print(aggregate(score~group+time,data,mean))
```
'''
    (fixture/"report.qmd").write_text(qmd)
    def quarto(directory):
        text=(directory/"report.html").read_text(encoding="utf-8")
        if "OBSERVATIONS=576 COMPLETE=535" not in text: raise AssertionError("Quarto executed result not found")
        return dict(executed_R=True,observations=576,complete_cases=535)
    run(job("quarto","render","quarto-executable-report",fixture/"report.qmd",{"format":"html"},data=prepared["data"]),quarto)
    tex='''\\documentclass{article}
\\usepackage{fontspec}
\\usepackage{xeCJK}
\\setCJKmainfont{Microsoft YaHei}
\\title{Research Bridge: 科研验收}
\\begin{document}
\\maketitle\\tableofcontents
\\section{Synthetic longitudinal study}\\label{sec:study}
144 subjects, 576 observations, 535 complete cases. 完整病例与缺失审计。
\\begin{equation} y=X\\beta+\\epsilon \\end{equation}
\\section{Cross-reference}Section~\\ref{sec:study} is the study definition.
\\end{document}
'''
    (fixture/"research.tex").write_text(tex,encoding="utf-8")
    def pdf(directory,name):
        from pypdf import PdfReader
        reader=PdfReader(directory/name)
        text="\n".join(p.extract_text() or "" for p in reader.pages)
        if "535" not in text or len(reader.pages)<1: raise AssertionError("PDF content oracle mismatch")
        return dict(pages=len(reader.pages),complete_cases=535)
    run(job("latex","render","xelatex-chinese-crossreferences",fixture/"research.tex"),lambda p:pdf(p,"research.pdf"))
    source_doc=outputs.get("pandoc-research-docx")
    if source_doc and (source_doc/"report.docx").exists():
        run(job("libreoffice","convert","libreoffice-document-pdf",source_doc/"report.docx",{"format":"pdf"}),lambda p:pdf(p,"report.pdf"))
        run(job("office","convert","word-document-pdf",source_doc/"report.docx",{"application":"word"}),lambda p:pdf(p,"report.pdf"))
    # The formula oracle is independent of both Excel and LibreOffice recalculation.
    try:
        import openpyxl
        book=openpyxl.Workbook(); sheet=book.active; sheet.title="Model"
        sheet.append(["Visit","Active A","Active B","Mean difference"])
        for i in range(1,5): sheet.append([i,50+2*i,50+3*i,f"=C{i+1}-B{i+1}"])
        sheet["F1"]="Mean contrast"; sheet["F2"]="=AVERAGE(D2:D5)"; sheet["F3"]="=SUM(D2:D5)"
        data_sheet=book.create_sheet("Study")
        study_rows=list(csv.DictReader(Path(prepared["data"]).open(encoding="utf-8")))
        data_sheet.append(["Subject","Group","Visit","Score"])
        for row in study_rows:
            data_sheet.append([int(row["subject"]),row["group"],int(row["time"]),float(row["score"]) if row["score"] else None])
        sheet["H1"]="Complete observations"; sheet["I1"]="=COUNT(Study!D2:D577)"
        sheet["H2"]="Missing observations"; sheet["I2"]="=COUNTBLANK(Study!D2:D577)"
        for index,group in enumerate(("control","active_a","active_b"),4):
            sheet[f"H{index}"]=group
            sheet[f"I{index}"]=f'=COUNTIF(Study!B2:B577,H{index})'
            sheet[f"J{index}"]=f'=AVERAGEIF(Study!B2:B577,H{index},Study!D2:D577)'
        book.save(fixture/"formula-model.xlsx")
        def formulas(directory):
            b=openpyxl.load_workbook(directory/"recalculated.xlsx",data_only=True)
            values=[b["Model"][c].value for c in ("D2","D3","D4","D5","F2","F3")]
            if values!=[1,2,3,4,2.5,10]: raise AssertionError("spreadsheet formula oracle mismatch: "+str(values))
            if [b["Model"][c].value for c in ("I1","I2")]!=[535,41]: raise AssertionError("spreadsheet missing-data audit mismatch")
            errors=[]
            for index,group in enumerate(("control","active_a","active_b"),4):
                data_values=[float(row["score"]) for row in study_rows if row["group"]==group and row["score"]]
                if b["Model"][f"I{index}"].value!=192: raise AssertionError("group sample count mismatch")
                errors.append(abs(b["Model"][f"J{index}"].value-sum(data_values)/len(data_values)))
            if max(errors)>1e-10: raise AssertionError("spreadsheet aggregate mean oracle mismatch")
            return dict(recalculated_values=values,observations=576,complete=535,missing=41,max_group_mean_abs_error=max(errors))
        run(job("office","recalculate","excel-formula-recalculation",fixture/"formula-model.xlsx",{"application":"excel"}),formulas)
    except ImportError:
        checks.append(dict(id="spreadsheet-fixture",engine="office",status="blocked",reason="openpyxl fixture dependency unavailable",application_verified=False))
    try:
        from pptx import Presentation
        deck=Presentation()
        for title,body in [("Synthetic longitudinal intervention","144 subjects; 576 observations; 535 complete cases"),("Research design","Control / Active A / Active B; 4 visits; random subject intercept"),("Transparent inference","OLS comparison; Holm multiplicity correction; 200 subject bootstrap replicates")]:
            slide=deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text=title; slide.placeholders[1].text=body
        deck.save(fixture/"study.pptx")
        def slides_pdf(directory):
            result=pdf(directory,"report.pdf")
            if result["pages"]!=3: raise AssertionError("PowerPoint slide/page count mismatch")
            return result
        run(job("office","convert","powerpoint-research-slides",fixture/"study.pptx",{"application":"powerpoint"}),slides_pdf)
    except ImportError:
        checks.append(dict(id="powerpoint-fixture",engine="office",status="blocked",reason="python-pptx fixture dependency unavailable",application_verified=False))
    notebook={"nbformat":4,"nbformat_minor":5,"metadata":{"kernelspec":{"display_name":"DSH Research Python","language":"python","name":"dsh-research"}},"cells":[{"cell_type":"code","execution_count":None,"id":"native-check","metadata":{},"outputs":[],"source":["import csv, pathlib\n","rows=list(csv.DictReader(open("+repr(prepared["data"])+")))\n","n=sum(bool(r['score']) for r in rows)\n","assert len(rows)==576 and n==535\n","print({'observations':len(rows),'complete':n})"]}]}
    notebook["cells"].append({"cell_type":"code","execution_count":None,"id":"scientific-model","metadata":{},"outputs":[],"source":["import pandas as pd, statsmodels.api as sm\n","data=pd.read_csv("+repr(prepared["data"])+").dropna(subset=['score'])\n","predictors="+repr(["time","treat_a","treat_b","time_a","time_b","baseline_c","age_c"])+"\n","fit=sm.OLS(data['score'], sm.add_constant(data[predictors])).fit()\n","expected="+repr(oracle["coefficients"])+"\n","actual={'Intercept':fit.params['const'], **fit.params.drop('const').to_dict()}\n","assert max(abs(actual[k]-v) for k,v in expected.items())<1e-7\n","print(fit.summary())\n","print(data.groupby(['group','time'])['score'].agg(['count','mean','std']))"]})
    save_json(fixture/"study.ipynb",notebook)
    def nb(directory):
        data=json.loads((directory/"executed.ipynb").read_text())
        for cell in data["cells"]:
            outputs=cell["outputs"]
            if not outputs or any(x.get("output_type")=="error" for x in outputs): raise AssertionError("notebook cell did not execute")
        return dict(executed_cells=2,observations=576,complete=535,statsmodels_OLS_oracle=True)
    run(job("jupyter","execute","jupyter-executed-notebook",fixture/"study.ipynb"),nb)
    run(job("git","snapshot","git-reproducible-snapshot",fixture))
    if args.zotero_fixture and (selected is None or "zotero" in selected):
        process=subprocess.run([sys.executable,"-B",str(Path(__file__).with_name("zotero_fixture.py")),str(root/"zotero-fixture")],capture_output=True,text=True,encoding="utf-8",timeout=110)
        evidence_path=root/"zotero-fixture/zotero-acceptance.json"
        if evidence_path.is_file():
            fixture_report=json.loads(evidence_path.read_text(encoding="utf-8"))
            checks.append(dict(id="zotero-isolated-library",engine="zotero",status=fixture_report["status"],scope="native_local_api",application_verified=fixture_report["status"]=="passed",oracle={k:fixture_report.get(k) for k in ("articles","limit_verified","unicode_and_doi","child_note")},evidence=str(evidence_path)))
        else:
            checks.append(dict(id="zotero-isolated-library",engine="zotero",status="failed",reason=process.stderr,application_verified=False))
    else:
        run(job("zotero","search","zotero-local-library-read",params={"query":"","limit":5}))

    counts=dict(Counter(c["status"] for c in checks))
    status="failed" if counts.get("failed") or not checks else ("passed" if counts.get("passed") else "blocked")
    result=dict(schema_version="research.acceptance.v1",status=status,date="2026-10-24",synthetic_data=True,scope="General statistics, numerical computing, literature, publication, figures, office and qualitative interchange",counts=counts,checks=checks,integrity=verify_directory(root/"runs"))
    save_json(root/"acceptance.json",result)
    print(json.dumps({"status":result["status"],"counts":counts,"report":str(root/"acceptance.json")}),flush=True)
    return 1 if result["status"]=="failed" else (2 if result["status"]=="blocked" else 0)

if __name__=="__main__": raise SystemExit(main())
